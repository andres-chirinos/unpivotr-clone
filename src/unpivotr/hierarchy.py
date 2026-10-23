"""Módulo de estructuración jerárquica para datos estadísticos multidimensionales tipo SDMX."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import openpyxl

from .coercers import to_slug
from .grid import Grid


@dataclass(frozen=True)
class HierarchyNode:
    """Nodo inmutable que representa un concepto dentro de una taxonomía o dimensión jerárquica SDMX."""

    row: int
    depth: int
    text: str
    slug: str
    parent_slug: str | None
    path_slug: str
    path_text: str


def parse_hierarchy_tree(
    source: str | Path | Grid,
    sheet_name: str | None = None,
    col_idx: int = 1,
    row_start: int = 1,
    row_end: int | None = None,
    section_break_fn: Callable[[int, Any], str | None] | None = None,
) -> dict[int, HierarchyNode]:
    """Extrae la jerarquía taxonómica de una columna considerando indentación por espacios o cell.alignment.indent.

    Produce identificadores dimensionales canónicos separados por puntos (estándar SDMX)
    y descripciones jerárquicas navegables ('Parent > Child > Subchild').
    """
    nodes: dict[int, HierarchyNode] = {}
    stack: list[tuple[int, str, str]] = []  # (indent, clean_text, slug)
    current_section: str | None = None

    if isinstance(source, (str, Path)):
        wb = openpyxl.load_workbook(source, read_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        try:
            for r, row_cells in enumerate(ws.iter_rows(), start=1):
                if r < row_start:
                    continue
                if row_end is not None and r > row_end:
                    break
                if col_idx - 1 >= len(row_cells):
                    continue
                cell = row_cells[col_idx - 1]
                val = cell.value
                if val is None:
                    continue
                s_val = str(val)
                if not s_val.strip() or s_val.lower().startswith("fuente"):
                    continue

                if section_break_fn:
                    new_sec = section_break_fn(r, val)
                    if new_sec is not None:
                        current_section = new_sec
                        stack = []
                        continue

                space_indent = len(s_val) - len(s_val.lstrip(" "))
                al = cell.alignment
                align_indent = int(al.indent) if (al and al.indent) else 0
                indent = space_indent if space_indent > 0 else align_indent

                clean_text = s_val.strip()
                slug = to_slug(clean_text).upper()

                while stack and stack[-1][0] >= indent:
                    stack.pop()

                parent_slugs = [item[2] for item in stack]
                parent_texts = [item[1] for item in stack]

                prefix_slug = [current_section] if current_section else []
                prefix_text = (
                    [current_section.replace("_", " ").title()] if current_section else []
                )

                all_slugs = prefix_slug + parent_slugs + [slug]
                all_texts = prefix_text + parent_texts + [clean_text]

                path_slug = ".".join(all_slugs)
                path_text = " > ".join(all_texts)
                parent_slug = (
                    parent_slugs[-1] if parent_slugs else (current_section or None)
                )

                depth = len(prefix_slug) + len(stack)
                stack.append((indent, clean_text, slug))

                nodes[r] = HierarchyNode(
                    row=r,
                    depth=depth,
                    text=clean_text,
                    slug=slug,
                    parent_slug=parent_slug,
                    path_slug=path_slug,
                    path_text=path_text,
                )
        finally:
            wb.close()
    elif isinstance(source, Grid):
        max_r = row_end if row_end is not None else source.max_row
        for r in range(row_start, max_r + 1):
            val = source.cell(r, col_idx).val
            if val is None:
                continue
            s_val = str(val)
            if not s_val.strip() or s_val.lower().startswith("fuente"):
                continue

            if section_break_fn:
                new_sec = section_break_fn(r, val)
                if new_sec is not None:
                    current_section = new_sec
                    stack = []
                    continue

            indent = len(s_val) - len(s_val.lstrip(" "))
            clean_text = s_val.strip()
            slug = to_slug(clean_text).upper()

            while stack and stack[-1][0] >= indent:
                stack.pop()

            parent_slugs = [item[2] for item in stack]
            parent_texts = [item[1] for item in stack]

            prefix_slug = [current_section] if current_section else []
            prefix_text = (
                [current_section.replace("_", " ").title()] if current_section else []
            )

            all_slugs = prefix_slug + parent_slugs + [slug]
            all_texts = prefix_text + parent_texts + [clean_text]

            path_slug = ".".join(all_slugs)
            path_text = " > ".join(all_texts)
            parent_slug = parent_slugs[-1] if parent_slugs else (current_section or None)

            depth = len(prefix_slug) + len(stack)
            stack.append((indent, clean_text, slug))

            nodes[r] = HierarchyNode(
                row=r,
                depth=depth,
                text=clean_text,
                slug=slug,
                parent_slug=parent_slug,
                path_slug=path_slug,
                path_text=path_text,
            )

    return nodes
