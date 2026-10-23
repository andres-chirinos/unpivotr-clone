"""Pruebas unitarias de invertibilidad matemática y auditoría de pérdida de información."""
import pytest
import pandas as pd
from unpivotr import (
    Grid,
    TableFlattener,
    InvertibilityError,
    InformationLossError,
    DataCorruptionError,
    InvertibilityReport,
    InvertibilityVerifier,
    verify_invertibility,
    synthesize_heterogeneous_workbook,
    to_slug,
)


def test_flattener_invert_and_verify_perfect_fidelity():
    """Comprueba que un flattener aplicado a una hoja no rectangular se invierte al 100%."""
    ground_truth = pd.DataFrame([
        {"id": "SERIE_A", "fecha": "2026-01-01", "valor": 100.5, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie A"},
        {"id": "SERIE_A", "fecha": "2026-01-02", "valor": 105.0, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie A"},
        {"id": "SERIE_B", "fecha": "2026-01-01", "valor": 200.0, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie B"},
        {"id": "SERIE_B", "fecha": "2026-01-02", "valor": 210.25, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie B"},
    ])

    wb = synthesize_heterogeneous_workbook(ground_truth, title="TABLA DE PRUEBA")
    ws = wb.active

    flattener = (
        TableFlattener("test_table")
        .from_sheet(ws)
        .data_bounds(row_start=8, col_start=2, row_end=9, col_end=3)
        .bind_date(6)
        .bind("id", lambda ctx: f"SERIE_{to_slug(ctx.col_val(1))[-1]}")
        .bind("descripcion", lambda ctx: str(ctx.col_val(1)))
        .filter(lambda ctx: not str(ctx.col_val(1)).startswith("TOTAL"))
    )

    df = flattener.to_dataframe()
    assert len(df) == 4

    # 1. Invertir a una malla reconstruida
    recon_grid = flattener.invert(df, preserve_headers=True)
    assert recon_grid.cell_value(8, 2) == 100.5
    assert recon_grid.cell_value(8, 3) == 105.0
    assert recon_grid.cell_value(9, 2) == 200.0
    assert recon_grid.cell_value(9, 3) == 210.25
    # Encabezados preservados
    assert recon_grid.cell_value(2, 1) == "TABLA DE PRUEBA"

    # 2. Verificación formal de invertibilidad
    report = flattener.verify_invertibility(df)
    assert isinstance(report, InvertibilityReport)
    assert report.is_invertible is True
    assert report.total_source_cells == 4
    assert report.total_extracted_observations == 4
    assert report.lost_cells_count == 0
    assert report.mismatched_cells_count == 0
    assert report.cell_coverage_pct == 100.0
    assert report.max_absolute_error == 0.0

    # No debe lanzar error
    report.assert_valid()


def test_invertibility_detects_lost_information():
    """Si la extracción ignora intencionalmente una columna, el verificador detecta la pérdida."""
    ground_truth = pd.DataFrame([
        {"id": "SERIE_A", "fecha": "2026-01-01", "valor": 100.5, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie A"},
        {"id": "SERIE_A", "fecha": "2026-01-02", "valor": 105.0, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie A"},
        {"id": "SERIE_B", "fecha": "2026-01-01", "valor": 200.0, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie B"},
        {"id": "SERIE_B", "fecha": "2026-01-02", "valor": 210.25, "categoria": "CAT", "unidad": "USD", "descripcion": "Serie B"},
    ])

    wb = synthesize_heterogeneous_workbook(ground_truth, title="TABLA RESTRINGIDA")
    ws = wb.active

    # Flattener con bounds truncados (solo col 2, perdiendo col 3)
    flattener = (
        TableFlattener("test_truncated")
        .from_sheet(ws)
        .data_bounds(row_start=8, col_start=2, row_end=9, col_end=2)
        .bind_date(6)
        .bind("id", lambda ctx: f"SERIE_{to_slug(ctx.col_val(1))[-1]}")
    )

    df = flattener.to_dataframe()
    assert len(df) == 2  # Solo 2 de las 4 observaciones

    # Verificamos contra los bounds reales de la hoja (col_end=3)
    verifier = InvertibilityVerifier()
    recon_grid = flattener.invert(df)
    report = verifier.verify_grids(
        source_grid=flattener._grid,
        reconstructed_grid=recon_grid,
        row_start=8,
        col_start=2,
        row_end=9,
        col_end=3,
    )

    assert report.is_invertible is False
    assert report.lost_cells_count == 2
    assert report.cell_coverage_pct == 50.0
    assert len(report.lost_cells) == 2
    assert report.lost_cells[0]["col"] == 3

    with pytest.raises(InformationLossError, match="Pérdida de información crítica"):
        report.assert_valid()


def test_invertibility_detects_value_corruption():
    """Si una transformación matemática corrompe los valores, se reporta discrepancia numérica."""
    source_grid = Grid.from_rows([
        ["Concepto", "2026-01-01", "2026-01-02"],
        ["Ventas", 100.0, 200.0],
    ])
    # Reconstrucción con un valor alterado
    recon_grid = Grid.from_rows([
        ["Concepto", "2026-01-01", "2026-01-02"],
        ["Ventas", 100.0, 205.5],  # Error de 5.5
    ])

    report = verify_invertibility(
        target=source_grid,
        reconstructed_or_df=recon_grid,
        row_start=2,
        col_start=2,
        row_end=2,
        col_end=3,
        tolerance=1e-4,
    )

    assert report.is_invertible is False
    assert report.mismatched_cells_count == 1
    assert report.mismatches[0]["diff"] == pytest.approx(5.5)
    assert report.max_absolute_error == pytest.approx(5.5)


def test_to_dataframe_include_coords():
    """Verifica que to_dataframe(include_coords=True) genera las columnas de procedencia."""
    grid = Grid.from_rows([
        ["Concepto", "2026-01-01"],
        ["Activo", 1500.0],
    ])
    flattener = (
        TableFlattener("coords_test")
        .from_grid(grid)
        .data_bounds(row_start=2, col_start=2, row_end=2, col_end=2)
        .bind_date(1)
        .bind_id(1)
    )

    df = flattener.to_dataframe(include_coords=True)
    assert "_row" in df.columns
    assert "_col" in df.columns
    assert df["_row"].iloc[0] == 2
    assert df["_col"].iloc[0] == 2


def test_matrix_formats_agnostic_invertibility():
    """Comprueba que verify_invertibility acepta numpy, DataFrame wide, listas y Grids indistintamente."""
    import numpy as np

    # 1. Matriz original como lista de listas
    source_matrix = [
        ["Concepto", "2026-01-01", "2026-01-02"],
        ["Serie A", 100.0, 200.0],
        ["Serie B", 300.0, 400.0],
    ]

    # 2. Reconstrucción como numpy array
    recon_np = np.array([
        ["Concepto", "2026-01-01", "2026-01-02"],
        ["Serie A", 100.0, 200.0],
        ["Serie B", 300.0, 400.0],
    ], dtype=object)

    report_np = verify_invertibility(source_matrix, recon_np, row_start=2, col_start=2)
    assert report_np.is_invertible is True
    assert report_np.total_source_cells == 4
    report_np.assert_valid()

    # 3. Reconstrucción como DataFrame tidy con _row y _col
    tidy_df = pd.DataFrame([
        {"id": "A", "fecha": "2026-01-01", "valor": 100.0, "_row": 2, "_col": 2},
        {"id": "A", "fecha": "2026-01-02", "valor": 200.0, "_row": 2, "_col": 3},
        {"id": "B", "fecha": "2026-01-01", "valor": 300.0, "_row": 3, "_col": 2},
        {"id": "B", "fecha": "2026-01-02", "valor": 400.0, "_row": 3, "_col": 3},
    ])
    report_tidy = verify_invertibility(source_matrix, tidy_df, row_start=2, col_start=2)
    assert report_tidy.is_invertible is True
    assert report_tidy.cell_coverage_pct == 100.0


def test_trivial_data_handling():
    """Celdas con guiones, 's.d.' o vacíos en el rango de datos no causan falso positivo de pérdida."""
    from unpivotr import is_trivial_value

    assert is_trivial_value("-") is True
    assert is_trivial_value("s.d.") is True
    assert is_trivial_value("n.d.") is True
    assert is_trivial_value("—") is True
    assert is_trivial_value(None) is True
    assert is_trivial_value(123.45) is False

    source_with_trivial = Grid.from_rows([
        ["Concepto", "2026-01-01", "2026-01-02"],
        ["Serie A", 100.0, "-"],        # Guión trivial
        ["Serie B", "s.d.", 200.0],    # s.d. trivial
    ])

    # Reconstrucción solo contiene las celdas numéricas reales
    recon_matrix = Grid.from_rows([
        ["Concepto", "2026-01-01", "2026-01-02"],
        ["Serie A", 100.0, None],
        ["Serie B", None, 200.0],
    ])

    report = verify_invertibility(source_with_trivial, recon_matrix, row_start=2, col_start=2, ignore_trivial=True)
    assert report.is_invertible is True
    assert report.total_source_cells == 2  # Solo 2 celdas numéricas reales
    assert report.trivial_cells_ignored == 2
    assert report.lost_cells_count == 0


def test_specialized_exceptions():
    """Verifica que assert_valid lance las excepciones correspondientes ante fallos."""
    from unpivotr import InformationLossError, DataCorruptionError

    # 1. Pérdida de datos
    source = Grid.from_rows([["A", 100.0], ["B", 200.0]])
    recon_loss = Grid.from_rows([["A", 100.0], ["B", None]])
    rep_loss = verify_invertibility(source, recon_loss)
    with pytest.raises(InformationLossError, match="Pérdida de información crítica"):
        rep_loss.assert_valid()

    # 2. Corrupción de valores
    recon_corrupt = Grid.from_rows([["A", 100.0], ["B", 250.0]])
    rep_corrupt = verify_invertibility(source, recon_corrupt)
    with pytest.raises(DataCorruptionError, match="Corrupción de datos numéricos"):
        rep_corrupt.assert_valid()


def test_base_table_converter_in_unpivotr(tmp_path):
    """Comprueba que BaseTableConverter funciona directamente desde unpivotr."""
    import openpyxl
    from unpivotr import BaseTableConverter

    # Crear archivo de prueba
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestSheet"
    ws["A1"] = "CUADRO TEST"
    ws["A3"] = "Concepto"
    ws["B3"] = "2026-01-01"
    ws["A5"] = "Serie 1"
    ws["B5"] = 999.0

    test_file = tmp_path / "test.xlsx"
    wb.save(test_file)

    class DummyConverter(BaseTableConverter):
        table_id = "dummy"
        sheet_name = "TestSheet"
        category = "CAT"
        unit = "USD"

        def extract(self, file_path):
            f = self.build_flattener(file_path).data_bounds(row_start=5, col_start=2).bind_date(3).bind_id(1)
            return f.to_dataframe()

    conv = DummyConverter()
    df = conv.extract(test_file)
    assert len(df) == 1
    assert df["valor"].iloc[0] == 999.0

    verif = conv.verify_reconstruction(df, test_file)
    assert verif["valid"] is True
    assert verif["total_observations"] == 1
    assert verif["invertibility"]["is_invertible"] is True
