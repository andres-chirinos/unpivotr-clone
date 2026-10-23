import re
import openpyxl
import pandas as pd
from unpivotr import TableFlattener, to_slug

MESES_NUM = {
    "ENERO": "01",
    "FEBRERO": "02",
    "MARZO": "03",
}

def test_multi_sheet_processing():
    # Crear un workbook con dos hojas mensuales: "Enero 2026" y "Febrero 2026"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Enero 2026"
    ws2 = wb.create_sheet("Febrero 2026")

    # Hoja 1: Enero 2026
    # Fila 1: Título
    # Fila 2: Cabecera: Día, Entidad, Monto
    # Fila 3: 15, "Banco Central", 500.0
    # Fila 4: 16, "Banco Central", 700.0
    ws1["A1"] = "MOVIMIENTOS DE EFECTIVO"
    ws1.append(["Dia", "Entidad", "Monto"])
    ws1.append([15, "Banco Central", 500.0])
    ws1.append([16, "Banco Central", 700.0])

    # Hoja 2: Febrero 2026
    ws2["A1"] = "MOVIMIENTOS DE EFECTIVO"
    ws2.append(["Dia", "Entidad", "Monto"])
    ws2.append([5, "Banco Central", 350.0])

    # Procesar ambas hojas iterativamente con TableFlattener
    all_dfs = []
    for sheet_name in wb.sheetnames:
        m = re.match(r"(?i)([a-z]+)\s+(\d{4})", sheet_name.strip())
        if not m:
            continue
        mes_txt, anio_txt = m.groups()
        mes_num = MESES_NUM[mes_txt.upper()]

        flattener = (
            TableFlattener(f"movimientos_{sheet_name}")
            .from_sheet(wb[sheet_name])
            .validate_anchor(cell="A1", pattern=r"(?i)movimientos\s+de\s+efectivo")
            .data_bounds(row_start=3, col_start=3)  # Monto en col 3
            .bind("fecha", lambda ctx, y=anio_txt, m=mes_num: f"{y}-{m}-{int(ctx.col_val(1)):02d}")
            .bind("id", lambda ctx: f"MOV_{to_slug(ctx.col_val(2))}")
            .bind("categoria", "OPERACIONES")
            .bind("unidad", "USD")
            .bind("descripcion", lambda ctx: f"Movimiento - {ctx.col_val(2)}")
        )
        all_dfs.append(flattener.to_dataframe())

    consolidated = pd.concat(all_dfs, ignore_index=True)
    assert len(consolidated) == 3
    fechas = sorted(consolidated["fecha"].tolist())
    assert fechas == ["2026-01-15", "2026-01-16", "2026-02-05"]
    assert all(v in [500.0, 700.0, 350.0] for v in consolidated["valor"].tolist())
