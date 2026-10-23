"""Pruebas unitarias para el analizador de taxonomías y jerarquías tipo SDMX."""

import openpyxl
from unpivotr import parse_hierarchy_tree, HierarchyNode, Grid

def test_parse_hierarchy_tree_spaces(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Taxonomia"

    ws["A1"] = "DETERMINANTES"
    ws["A2"] = "  Reservas Internacionales"
    ws["A3"] = "    Divisas"
    ws["A4"] = "    Oro"
    ws["A5"] = "COMPONENTES"
    ws["A6"] = "  Billetes y Monedas"

    fpath = tmp_path / "test_spaces.xlsx"
    wb.save(fpath)

    nodes = parse_hierarchy_tree(fpath, sheet_name="Taxonomia", row_start=1)
    assert len(nodes) == 6
    assert nodes[1].path_slug == "DETERMINANTES"
    assert nodes[1].depth == 0
    assert nodes[2].path_slug == "DETERMINANTES.RESERVAS_INTERNACIONALES"
    assert nodes[2].depth == 1
    assert nodes[3].path_slug == "DETERMINANTES.RESERVAS_INTERNACIONALES.DIVISAS"
    assert nodes[3].depth == 2
    assert nodes[3].path_text == "DETERMINANTES > Reservas Internacionales > Divisas"
    assert nodes[3].parent_slug == "RESERVAS_INTERNACIONALES"
    assert nodes[5].path_slug == "COMPONENTES"
    assert nodes[5].depth == 0
    assert nodes[6].path_slug == "COMPONENTES.BILLETES_Y_MONEDAS"
    assert nodes[6].depth == 1

def test_parse_hierarchy_from_grid():
    rows = [
        ["SECTOR PUBLICO"],
        ["  GOBIERNO CENTRAL"],
        ["    MINISTERIO"],
    ]
    grid = Grid.from_rows(rows)
    nodes = parse_hierarchy_tree(grid, row_start=1)
    assert len(nodes) == 3
    assert nodes[1].path_slug == "SECTOR_PUBLICO"
    assert nodes[2].path_slug == "SECTOR_PUBLICO.GOBIERNO_CENTRAL"
    assert nodes[3].path_slug == "SECTOR_PUBLICO.GOBIERNO_CENTRAL.MINISTERIO"
